import json
import tempfile
import unittest
from pathlib import Path

from main import ParseReport, Question, XXT


class CrawlClient(XXT):
    def get_schools(self):
        return [("school", "测试学校")]

    def get_courses(self, _school_id):
        return [("course", "测试课程")]

    def get_test_ids(self, *_args):
        return ["directory"]

    def get_exam_id(self, *_args):
        return "exam"

    def get_answer_url(self, *_args):
        return "answer-url"

    def get_question_report(self, _url):
        return ParseReport([
            Question(1, "选择题", "正确答案：A", options={"A": "答案 A", "B": "答案 B"}),
            Question(2, "填空题", "", "missing_answer"),
        ], ["第 2 题缺少答案节点"])


class CrawlTests(unittest.TestCase):
    def test_manual_parameters_do_not_depend_on_course_discovery(self):
        client = CrawlClient()
        client.get_schools = lambda: self.fail("manual recovery must not need the course list")

        with tempfile.TemporaryDirectory() as directory:
            report = client.crawl_course(
                "course", class_id="class", cpi="cpi", output_dir=directory, formats=("json",)
            )

        self.assertEqual("unknown_course", report.course_name)

    def test_manual_course_parameters_export_a_report_without_requesting_mode_page(self):
        with tempfile.TemporaryDirectory() as directory:
            report = CrawlClient().crawl_course(
                "course",
                class_id="class",
                cpi="cpi",
                output_dir=directory,
                formats=("json",),
            )

            self.assertEqual(2, report.total_questions)
            self.assertEqual(1, report.complete_questions)
            self.assertEqual(1, report.incomplete_questions)
            output = Path(directory) / "unknown_course(course)-latest.json"
            payload = json.loads(output.read_text(encoding="utf8"))
            self.assertEqual("选择题", payload[0]["question"])
            self.assertEqual({"A": "答案 A", "B": "答案 B"}, payload[0]["options"])
            self.assertEqual("missing_answer", payload[1]["status"])

    def test_xlsx_export_includes_parse_status(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            CrawlClient().crawl_course(
                "course", class_id="class", cpi="cpi", output_dir=directory, formats=("xlsx",)
            )

            workbook = load_workbook(Path(directory) / "unknown_course(course)-latest.xlsx")
            rows = list(workbook.active.iter_rows(values_only=True))
            self.assertEqual(("题号", "题目", "A", "B", "C", "D", "E", "F", "答案", "解析状态"), rows[0])
            self.assertEqual((2, "填空题", None, None, None, None, None, None, None, "missing_answer"), rows[2])


if __name__ == "__main__":
    unittest.main()
